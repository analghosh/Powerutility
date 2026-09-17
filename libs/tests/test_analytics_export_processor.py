"""The scheduler worker that turns a queued export job into a file on disk.

This is the first job in EPMS that writes a file, so most of what these tests pin
down has no precedent to copy and would otherwise only be discovered in
production, at 2am, on the single EC2 box that also serves the bidding screens:

* **Chunks must tile the date range exactly.** A gap loses days silently — the file
  looks complete and is short. An overlap double-counts them. Both are invisible in
  a 200,000-row spreadsheet, so the tiling is asserted as an identity, not a count.
  Chunk boundaries also have to respect the date grouping: splitting a month in two
  produces two half-month averages that no assembly step can put back together.

* **A retry must resume, not restart.** The whole point of writing an ordered part
  file per chunk is that attempt two starts at chunk 4 of 11. If the worker
  re-runs the first three, a flaky chunk turns a four-minute export into a job that
  never finishes.

* **The worker must stop when it is told to.** Two independent ways exist: the user
  cancels, and the reaper requeues the job away from this worker (``heartbeat``
  returns ``False``). A worker that keeps writing parts after either one is writing
  into a job another worker now owns.

* **Expiry deletes the file and keeps the row.** The row carries the filters that
  power Regenerate for the remaining 90 days. Deleting it is how a user loses the
  34 plants they picked three days ago.

Mongo is a small in-memory fake rather than an ``AsyncMock`` here, because half of
these behaviours *are* the query — "claims queued, ignores a running job with a
fresh heartbeat" is a statement about a filter, and asserting it on ``call_args``
would only restate the code. The fake speaks the subset of the query language the
repository actually uses and raises on anything else, so it cannot quietly agree
with a query it does not understand.
"""
import ast
import copy
import os
import sys
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from services.processing.app.v2.repository import (  # noqa: E402
    analytics_job_repository as jobs,
)
from services.scheduler.app.processors import (  # noqa: E402
    analytics_export_processor as worker,
)

REPO_ROOT = Path(__file__).resolve().parents[2]
CRON_MANAGER = REPO_ROOT / "services/scheduler/app/services/cron_manager.py"
JOB_SEEDER = REPO_ROOT / "services/scheduler/app/seeder/job_seeder.py"

NOW = datetime(2026, 9, 17, 10, 30, 0, tzinfo=timezone.utc)


# --------------------------------------------------------------------------
# A Mongo stand-in that understands exactly the operators the repository uses.
# --------------------------------------------------------------------------

class _Result:
    def __init__(self, matched, modified):
        self.matched_count = matched
        self.modified_count = modified


def _field(doc, path):
    """Dotted field read, with ``parts.idx`` returning the array's values."""
    current = doc
    for i, segment in enumerate(path.split(".")):
        if isinstance(current, list):
            return [(item or {}).get(segment) for item in current]
        if not isinstance(current, dict):
            return None
        current = current.get(segment)
        if current is None and i < len(path.split(".")) - 1:
            return None
    return current


def _resolve(doc, operand):
    if isinstance(operand, str) and operand.startswith("$"):
        return _field(doc, operand[1:])
    return operand


def _expr(doc, expression):
    for op, operands in expression.items():
        left, right = (_resolve(doc, o) for o in operands)
        if left is None or right is None:
            return False
        if op == "$lt" and not left < right:
            return False
        if op == "$gte" and not left >= right:
            return False
        if op not in ("$lt", "$gte"):
            raise AssertionError(f"fake collection cannot evaluate $expr {op}")
    return True


def _match_condition(value, cond):
    for op, operand in cond.items():
        if op == "$in":
            if value not in operand:
                return False
        elif op == "$nin":
            if value in operand:
                return False
        elif op == "$ne":
            # On an array field (parts.idx) $ne means "no element equals".
            if isinstance(value, list):
                if operand in value:
                    return False
            elif value == operand:
                return False
        elif op in ("$lt", "$lte", "$gt", "$gte"):
            if value is None:
                return False
            if op == "$lt" and not value < operand:
                return False
            if op == "$lte" and not value <= operand:
                return False
            if op == "$gt" and not value > operand:
                return False
            if op == "$gte" and not value >= operand:
                return False
        else:
            raise AssertionError(f"fake collection cannot evaluate {op}")
    return True


def _matches(doc, query):
    for key, cond in (query or {}).items():
        if key == "$expr":
            if not _expr(doc, cond):
                return False
            continue
        if key == "$or":
            if not any(_matches(doc, q) for q in cond):
                return False
            continue
        value = _field(doc, key)
        if isinstance(cond, dict) and any(k.startswith("$") for k in cond):
            if not _match_condition(value, cond):
                return False
        elif value != cond:
            return False
    return True


def _apply(doc, update):
    for op, fields in update.items():
        if op == "$set":
            doc.update(fields)
        elif op == "$inc":
            for key, delta in fields.items():
                doc[key] = (doc.get(key) or 0) + delta
        elif op == "$push":
            for key, item in fields.items():
                doc.setdefault(key, []).append(item)
        elif op == "$pull":
            for key, cond in fields.items():
                doc[key] = [item for item in (doc.get(key) or [])
                            if not _matches(item, cond)]
        else:
            raise AssertionError(f"fake collection cannot apply {op}")


def _project(doc, projection):
    out = copy.deepcopy(doc)
    if projection:
        for key, keep in projection.items():
            if not keep:
                out.pop(key, None)
    return out


class FakeCollection:
    """Enough MongoDB to run the real repository against."""

    def __init__(self, docs=None):
        self.docs = [copy.deepcopy(d) for d in (docs or [])]
        self.deletes = 0

    def _sorted(self, candidates, sort):
        if not sort:
            return candidates
        for key, direction in reversed(sort):
            candidates.sort(
                key=lambda d: (d.get(key) is None, d.get(key)),
                reverse=direction < 0,
            )
        return candidates

    async def insert_one(self, doc):
        self.docs.append(copy.deepcopy(doc))
        return _Result(1, 1)

    async def find_one(self, query, projection=None):
        for doc in self.docs:
            if _matches(doc, query):
                return _project(doc, projection)
        return None

    async def update_one(self, query, update):
        for doc in self.docs:
            if _matches(doc, query):
                _apply(doc, update)
                return _Result(1, 1)
        return _Result(0, 0)

    async def update_many(self, query, update):
        hit = 0
        for doc in self.docs:
            if _matches(doc, query):
                _apply(doc, update)
                hit += 1
        return _Result(hit, hit)

    async def find_one_and_update(self, query, update, projection=None,
                                  sort=None, return_document=None):
        for doc in self._sorted([d for d in self.docs if _matches(d, query)], sort):
            _apply(doc, update)
            return _project(doc, projection)
        return None

    async def delete_one(self, query):  # pragma: no cover - must never be called
        self.deletes += 1
        raise AssertionError("the export job row must outlive its file")

    def by_id(self, job_id):
        for doc in self.docs:
            if doc.get("job_id") == job_id:
                return doc
        return None


# --------------------------------------------------------------------------
# Fixtures
# --------------------------------------------------------------------------

def make_filters(**overrides):
    filters = {
        "datasets": ["ACTUAL_DEMAND"],
        "from_date": "2026-01-01",
        "to_date": "2026-04-10",          # 100 days
        "from_block": 1,
        "to_block": 96,
        "date_grouping": "BLOCK_WISE",
        "block_grouping": "INDIVIDUAL",
        "plants": [],
        "format": "csv",
    }
    filters.update(overrides)
    return filters


def make_job(**overrides):
    job = {
        "job_id": "job-1",
        "user_id": "u-1",
        "request_hash": "h-1",
        "filters": make_filters(),
        "label": "Actual Demand - 100 days",
        "status": jobs.STATUS_QUEUED,
        "cells_estimate": 500_000,
        "chunks_total": 0,
        "chunks_done": 0,
        "parts": [],
        "result": None,
        "error": None,
        "attempts": 0,
        "max_attempts": 3,
        "worker_id": None,
        "heartbeat_at": None,
        "created_at": NOW - timedelta(minutes=5),
        "started_at": None,
        "finished_at": None,
        "expires_at": None,
    }
    filters = overrides.pop("filters", None)
    if filters is not None:
        job["filters"] = filters
    job.update(overrides)
    return job


class RecordingSource:
    """Stands in for the aggregation. Records the window each chunk asked for."""

    def __init__(self, rows_per_chunk=2, on_call=None):
        self.windows = []
        self.rows_per_chunk = rows_per_chunk
        self.on_call = on_call

    async def __call__(self, filters, datasets):
        self.windows.append((filters["from_date"], filters["to_date"]))
        if self.on_call is not None:
            await self.on_call(self, filters)
        rows = [
            {"dataset": datasets[0], "date_key": filters["from_date"],
             "block_key": str(i + 1), "entity": None, "value": float(i)}
            for i in range(self.rows_per_chunk)
        ]
        return rows, []


def build(tmp_path, docs, *, source=None, **kwargs):
    """A processor wired to a fake Mongo and a temp export directory."""
    repo = jobs.AnalyticsJobRepository()
    collection = FakeCollection(docs)
    repo.collection = collection
    processor = worker.AnalyticsExportProcessor(
        job_repository=repo,
        query_runner=source if source is not None else RecordingSource(),
        export_dir=str(tmp_path / "exports"),
        worker_id="worker-A",
        **kwargs,
    )
    return processor, collection


# --------------------------------------------------------------------------
# chunks_for - the tiling
# --------------------------------------------------------------------------

def test_a_500k_cell_job_splits_into_about_ten_chunks_of_about_50k(tmp_path):
    processor, _ = build(tmp_path, [])

    chunks = processor.chunks_for(make_job())

    assert len(chunks) == 10
    assert {c["cells"] for c in chunks} == {50_000}


def test_chunks_tile_the_date_range_with_no_gap_and_no_overlap(tmp_path):
    processor, _ = build(tmp_path, [])
    job = make_job()

    chunks = processor.chunks_for(job)

    start = date.fromisoformat(job["filters"]["from_date"])
    end = date.fromisoformat(job["filters"]["to_date"])

    # The union of the chunks is the original range, day for day, once each.
    covered = []
    for chunk in chunks:
        day = date.fromisoformat(chunk["from_date"])
        last = date.fromisoformat(chunk["to_date"])
        assert day <= last
        while day <= last:
            covered.append(day)
            day += timedelta(days=1)

    expected = [start + timedelta(days=i) for i in range((end - start).days + 1)]
    assert covered == expected


def test_chunk_indexes_are_contiguous_from_zero(tmp_path):
    processor, _ = build(tmp_path, [])

    chunks = processor.chunks_for(make_job())

    assert [c["idx"] for c in chunks] == list(range(len(chunks)))


def test_a_small_job_is_a_single_chunk(tmp_path):
    processor, _ = build(tmp_path, [])
    job = make_job(cells_estimate=1_000,
                   filters=make_filters(from_date="2026-01-01", to_date="2026-01-03"))

    chunks = processor.chunks_for(job)

    assert len(chunks) == 1
    assert chunks[0]["from_date"] == "2026-01-01"
    assert chunks[0]["to_date"] == "2026-01-03"


def test_month_wise_chunks_never_split_a_month(tmp_path):
    """Half a month aggregated on its own is a half-month average, and no assembly
    step can turn two of those back into the month the user asked for."""
    processor, _ = build(tmp_path, [])
    job = make_job(
        cells_estimate=5_000_000,   # forces the smallest possible chunk
        filters=make_filters(from_date="2026-01-05", to_date="2026-04-10",
                             date_grouping="MONTH_WISE"),
    )

    chunks = processor.chunks_for(job)

    for chunk in chunks[:-1]:
        last = date.fromisoformat(chunk["to_date"])
        # Every interior boundary lands on a month end.
        assert (last + timedelta(days=1)).day == 1
    assert chunks[0]["from_date"] == "2026-01-05"
    assert chunks[-1]["to_date"] == "2026-04-10"


def test_week_wise_chunks_never_split_an_iso_week(tmp_path):
    processor, _ = build(tmp_path, [])
    job = make_job(
        cells_estimate=5_000_000,
        filters=make_filters(from_date="2026-01-01", to_date="2026-02-28",
                             date_grouping="WEEK_WISE"),
    )

    chunks = processor.chunks_for(job)

    for chunk in chunks[:-1]:
        # ISO weeks end on Sunday (weekday 6).
        assert date.fromisoformat(chunk["to_date"]).weekday() == 6


# --------------------------------------------------------------------------
# run_job - resume, cancellation, lost claim
# --------------------------------------------------------------------------

def seed_parts(processor, job, count):
    """Real part files for the first `count` chunks, written by the worker's own
    writer so the test cannot drift from the format assembly expects."""
    parts_dir = os.path.join(processor.export_dir, job["job_id"], "parts")
    os.makedirs(parts_dir, exist_ok=True)
    receipts = []
    for chunk in processor.chunks_for(job)[:count]:
        rows = [{"dataset": job["filters"]["datasets"][0],
                 "date_key": chunk["from_date"], "block_key": "1",
                 "entity": None, "value": 1.0}]
        receipts.append(processor._write_part(parts_dir, chunk, rows, []))
    return receipts


@pytest.mark.asyncio
async def test_a_job_with_four_part_receipts_resumes_at_chunk_four(tmp_path):
    source = RecordingSource()
    processor, collection = build(tmp_path, [], source=source)
    job = make_job()
    job["parts"] = seed_parts(processor, job, 4)
    job["chunks_done"] = 4
    collection.docs.append(copy.deepcopy(job))

    result = await processor.run_job("job-1")

    ran = sorted(w[0] for w in source.windows)
    chunks = processor.chunks_for(job)
    assert ran == sorted(c["from_date"] for c in chunks[4:])
    assert len(ran) == 6

    # The four chunks it skipped are still in the finished file.
    assert result["status"] == jobs.STATUS_READY
    text = Path(collection.by_id("job-1")["result"]["path"]).read_text(encoding="utf-8")
    assert source_dates(text) == {c["from_date"] for c in chunks}


@pytest.mark.asyncio
async def test_a_receipt_whose_part_file_vanished_is_run_again(tmp_path):
    """Something cleared the working directory under a running job. Re-run those
    chunks; failing at assembly on a file that is not there is the worse answer."""
    source = RecordingSource()
    processor, collection = build(tmp_path, [], source=source)
    job = make_job()
    job["parts"] = seed_parts(processor, job, 4)
    job["chunks_done"] = 4
    collection.docs.append(copy.deepcopy(job))

    os.unlink(job["parts"][1]["path"])

    result = await processor.run_job("job-1")

    assert result["status"] == jobs.STATUS_READY
    chunks = processor.chunks_for(job)
    # Chunk 1 is back in the run, and its receipt was replaced rather than doubled.
    assert chunks[1]["from_date"] in {w[0] for w in source.windows}
    assert len(source.windows) == 7
    receipts = collection.by_id("job-1")["parts"]
    assert sorted(p["idx"] for p in receipts) == list(range(10))


@pytest.mark.asyncio
async def test_cancellation_is_observed_between_chunks(tmp_path):
    async def cancel_after_first(source, _filters):
        if len(source.windows) == 1:
            doc = collection.by_id("job-1")
            doc["status"] = jobs.STATUS_CANCELLED

    source = RecordingSource(on_call=cancel_after_first)
    processor, collection = build(tmp_path, [make_job()], source=source,
                                  chunk_concurrency=1)

    result = await processor.run_job("job-1")

    assert result["status"] == jobs.STATUS_CANCELLED
    assert len(source.windows) == 1, "the worker kept going after the cancel"
    assert collection.by_id("job-1")["status"] == jobs.STATUS_CANCELLED
    assert collection.by_id("job-1")["result"] is None


@pytest.mark.asyncio
async def test_a_cancel_during_the_last_chunk_still_stops_the_export(tmp_path):
    """There is no chunk after the last one, so without a check before assembly a
    late cancel would still hand the user a file."""
    async def cancel_on_last(source, _filters):
        if len(source.windows) == 10:
            collection.by_id("job-1")["status"] = jobs.STATUS_CANCELLED

    source = RecordingSource(on_call=cancel_on_last)
    processor, collection = build(tmp_path, [make_job()], source=source,
                                  chunk_concurrency=1)

    result = await processor.run_job("job-1")

    assert len(source.windows) == 10, "every chunk should have run"
    assert result["status"] == jobs.STATUS_CANCELLED
    assert collection.by_id("job-1")["result"] is None


@pytest.mark.asyncio
async def test_a_fully_resumed_job_assembles_without_running_a_chunk(tmp_path):
    """A job requeued right after its last chunk has nothing left to query — it
    still has to produce the file."""
    source = RecordingSource()
    processor, collection = build(tmp_path, [], source=source)
    job = make_job()
    job["parts"] = seed_parts(processor, job, 10)
    job["chunks_done"] = 10
    collection.docs.append(copy.deepcopy(job))

    result = await processor.run_job("job-1")

    assert source.windows == []
    assert result["status"] == jobs.STATUS_READY
    text = Path(collection.by_id("job-1")["result"]["path"]).read_text(encoding="utf-8")
    assert len(source_dates(text)) == 10


@pytest.mark.asyncio
async def test_a_worker_that_loses_its_claim_stops_at_the_next_chunk(tmp_path):
    """heartbeat() returning False means the reaper requeued the job away from this
    worker and someone else owns it now. Continuing writes parts into their job."""
    async def steal_after_first(source, _filters):
        if len(source.windows) == 1:
            collection.by_id("job-1")["worker_id"] = "worker-B"

    source = RecordingSource(on_call=steal_after_first)
    processor, collection = build(tmp_path, [make_job()], source=source,
                                  chunk_concurrency=1)

    result = await processor.run_job("job-1")

    assert result["status"] == worker.OUTCOME_LOST_CLAIM
    assert len(source.windows) == 1
    # The thief's claim is untouched and the job was not marked failed.
    assert collection.by_id("job-1")["worker_id"] == "worker-B"
    assert collection.by_id("job-1")["status"] == jobs.STATUS_RUNNING


@pytest.mark.asyncio
async def test_run_job_writes_a_csv_and_marks_the_job_ready(tmp_path):
    processor, collection = build(tmp_path, [make_job()])

    result = await processor.run_job("job-1")

    assert result["status"] == jobs.STATUS_READY
    doc = collection.by_id("job-1")
    assert doc["status"] == jobs.STATUS_READY
    written = doc["result"]
    assert written["filename"].endswith(".csv")
    assert os.path.exists(written["path"])
    assert written["bytes"] > 0
    assert written["row_count"] > 0

    text = Path(written["path"]).read_text(encoding="utf-8")
    # One filter header block and one column header for the whole file, not one
    # of each per chunk.
    assert text.count("EPMS Analytics Export") == 1
    assert text.count("dataset,label,entity,unit,date,block,value") == 1
    # Every chunk's rows made it in.
    assert len(source_dates(text)) == 10


def source_dates(text):
    dates = set()
    for line in text.splitlines():
        if line.startswith("#") or line.startswith("dataset,"):
            continue
        parts = line.split(",")
        if len(parts) >= 5 and parts[4]:
            dates.add(parts[4])
    return dates


@pytest.mark.asyncio
async def test_part_files_are_deleted_once_the_export_is_assembled(tmp_path):
    processor, collection = build(tmp_path, [make_job()])

    await processor.run_job("job-1")

    doc = collection.by_id("job-1")
    assert doc["parts"], "the receipts stay on the row"
    for part in doc["parts"]:
        assert not os.path.exists(part["path"])
    assert os.path.exists(doc["result"]["path"])


@pytest.mark.asyncio
async def test_the_export_directory_is_created_when_it_does_not_exist(tmp_path):
    target = tmp_path / "does" / "not" / "exist"
    processor, collection = build(tmp_path, [make_job()])
    processor.export_dir = str(target)

    await processor.run_job("job-1")

    assert target.is_dir()
    assert collection.by_id("job-1")["status"] == jobs.STATUS_READY


@pytest.mark.asyncio
async def test_an_unwritable_export_directory_fails_the_job_with_a_clear_error(tmp_path):
    """The PROD default is /var/lib/epms/exports, which an unprivileged service user
    cannot create. That must read as a configuration error on the job row, not as a
    bare PermissionError in a log nobody is watching."""
    locked = tmp_path / "locked"
    locked.mkdir()
    os.chmod(locked, 0o500)
    try:
        processor, collection = build(tmp_path, [make_job()])
        processor.export_dir = str(locked / "exports")

        result = await processor.run_job("job-1")

        assert result["status"] == jobs.STATUS_FAILED
        doc = collection.by_id("job-1")
        assert doc["status"] == jobs.STATUS_FAILED
        message = doc["error"]["message"]
        assert "ANALYTICS_EXPORT_DIR" in message
        assert str(locked) in message
    finally:
        os.chmod(locked, 0o700)


@pytest.mark.asyncio
async def test_a_cancelled_job_is_never_started(tmp_path):
    source = RecordingSource()
    processor, collection = build(
        tmp_path, [make_job(status=jobs.STATUS_CANCELLED)], source=source)

    result = await processor.run_job("job-1")

    assert result["status"] == worker.OUTCOME_SKIPPED
    assert source.windows == []


# --------------------------------------------------------------------------
# sweep
# --------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_sweep_claims_a_queued_job_and_leaves_a_live_running_one_alone(tmp_path):
    queued = make_job(job_id="job-queued")
    running = make_job(job_id="job-running", status=jobs.STATUS_RUNNING,
                       worker_id="worker-B", heartbeat_at=NOW - timedelta(seconds=5))
    processor, collection = build(tmp_path, [queued, running])
    processor.now = lambda: NOW

    result = await processor.sweep()

    assert result["claimed"] == ["job-queued"]
    assert collection.by_id("job-queued")["status"] == jobs.STATUS_READY
    # Untouched: still running, still owned by the other worker.
    other = collection.by_id("job-running")
    assert other["status"] == jobs.STATUS_RUNNING
    assert other["worker_id"] == "worker-B"


@pytest.mark.asyncio
async def test_sweep_claims_the_oldest_queued_job_first(tmp_path):
    older = make_job(job_id="job-old", created_at=NOW - timedelta(minutes=30))
    newer = make_job(job_id="job-new", created_at=NOW - timedelta(minutes=1))
    processor, _ = build(tmp_path, [newer, older], max_jobs_per_sweep=1)
    processor.now = lambda: NOW

    result = await processor.sweep()

    assert result["claimed"] == ["job-old"]


@pytest.mark.asyncio
async def test_sweep_stamps_the_worker_id_so_the_reaper_can_tell_claims_apart(tmp_path):
    processor, collection = build(tmp_path, [make_job()])
    processor.now = lambda: NOW

    await processor.sweep()

    assert collection.by_id("job-1")["worker_id"] == "worker-A"


# --------------------------------------------------------------------------
# reap
# --------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_reap_requeues_a_stale_running_job_and_increments_attempts(tmp_path):
    stale = make_job(status=jobs.STATUS_RUNNING, worker_id="worker-dead",
                     attempts=0, heartbeat_at=NOW - timedelta(minutes=10))
    processor, collection = build(tmp_path, [stale])
    processor.now = lambda: NOW

    result = await processor.reap()

    assert result["requeued"] == 1
    doc = collection.by_id("job-1")
    assert doc["status"] == jobs.STATUS_QUEUED
    assert doc["attempts"] == 1
    assert doc["worker_id"] is None


@pytest.mark.asyncio
async def test_reap_leaves_a_running_job_with_a_fresh_heartbeat_alone(tmp_path):
    live = make_job(status=jobs.STATUS_RUNNING, worker_id="worker-B",
                    heartbeat_at=NOW - timedelta(seconds=10))
    processor, collection = build(tmp_path, [live])
    processor.now = lambda: NOW

    result = await processor.reap()

    assert result["requeued"] == 0
    assert collection.by_id("job-1")["status"] == jobs.STATUS_RUNNING
    assert collection.by_id("job-1")["attempts"] == 0


@pytest.mark.asyncio
async def test_a_job_at_max_attempts_fails_instead_of_being_requeued(tmp_path):
    exhausted = make_job(status=jobs.STATUS_RUNNING, worker_id="worker-dead",
                         attempts=3, max_attempts=3,
                         heartbeat_at=NOW - timedelta(minutes=10))
    processor, collection = build(tmp_path, [exhausted])
    processor.now = lambda: NOW

    result = await processor.reap()

    assert result["requeued"] == 0
    doc = collection.by_id("job-1")
    assert doc["status"] == jobs.STATUS_FAILED
    assert doc["attempts"] == 3
    assert doc["error"]["code"] == "max_attempts_exceeded"


@pytest.mark.asyncio
async def test_reap_rescues_a_job_abandoned_during_assembly(tmp_path):
    """claim_orphans only looks at `running`. A worker killed while writing the file
    leaves its job in `assembling`, which nothing else would ever move — and
    find_active_by_hash counts that as in-flight, so the user cannot re-request it."""
    stuck = make_job(status=jobs.STATUS_ASSEMBLING, worker_id="worker-dead",
                     attempts=0, heartbeat_at=NOW - timedelta(minutes=10))
    processor, collection = build(tmp_path, [stuck])
    processor.now = lambda: NOW

    result = await processor.reap()

    assert result["requeued"] == 1
    doc = collection.by_id("job-1")
    assert doc["status"] == jobs.STATUS_QUEUED
    assert doc["attempts"] == 1
    assert doc["worker_id"] is None


@pytest.mark.asyncio
async def test_a_live_assembly_is_left_alone(tmp_path):
    live = make_job(status=jobs.STATUS_ASSEMBLING, worker_id="worker-B",
                    heartbeat_at=NOW - timedelta(seconds=10))
    processor, collection = build(tmp_path, [live])
    processor.now = lambda: NOW

    assert (await processor.reap())["requeued"] == 0
    assert collection.by_id("job-1")["status"] == jobs.STATUS_ASSEMBLING


@pytest.mark.asyncio
async def test_an_assembly_at_max_attempts_fails_rather_than_cycling(tmp_path):
    exhausted = make_job(status=jobs.STATUS_ASSEMBLING, worker_id="worker-dead",
                         attempts=3, max_attempts=3,
                         heartbeat_at=NOW - timedelta(minutes=10))
    processor, collection = build(tmp_path, [exhausted])
    processor.now = lambda: NOW

    result = await processor.reap()

    assert result["requeued"] == 0
    doc = collection.by_id("job-1")
    assert doc["status"] == jobs.STATUS_FAILED
    assert doc["error"]["code"] == "max_attempts_exceeded"


@pytest.mark.asyncio
async def test_reap_deletes_an_expired_file_sets_expired_and_keeps_the_row(tmp_path):
    stored = tmp_path / "exports" / "job-1" / "analytics.csv"
    stored.parent.mkdir(parents=True)
    stored.write_text("data", encoding="utf-8")

    ready = make_job(
        status=jobs.STATUS_READY,
        expires_at=NOW - timedelta(hours=1),
        download_token="tok",
        result={"path": str(stored), "filename": "analytics.csv",
                "bytes": 4, "row_count": 1},
    )
    processor, collection = build(tmp_path, [ready])
    processor.now = lambda: NOW

    result = await processor.reap()

    assert result["expired"] == 1
    assert result["files_deleted"] == 1
    assert not stored.exists()

    doc = collection.by_id("job-1")
    assert doc is not None, "the row carries the filters that power Regenerate"
    assert doc["status"] == jobs.STATUS_EXPIRED
    assert doc["filters"]["plants"] == []
    assert doc["download_token"] is None
    assert collection.deletes == 0


@pytest.mark.asyncio
async def test_reap_does_not_touch_a_file_that_is_still_inside_its_retention(tmp_path):
    stored = tmp_path / "exports" / "job-1" / "analytics.csv"
    stored.parent.mkdir(parents=True)
    stored.write_text("data", encoding="utf-8")

    ready = make_job(status=jobs.STATUS_READY,
                     expires_at=NOW + timedelta(hours=10),
                     result={"path": str(stored), "filename": "analytics.csv",
                             "bytes": 4, "row_count": 1})
    processor, collection = build(tmp_path, [ready])
    processor.now = lambda: NOW

    result = await processor.reap()

    assert result["expired"] == 0
    assert stored.exists()
    assert collection.by_id("job-1")["status"] == jobs.STATUS_READY


@pytest.mark.asyncio
async def test_an_already_missing_file_still_expires_the_row(tmp_path):
    ready = make_job(status=jobs.STATUS_READY,
                     expires_at=NOW - timedelta(hours=1),
                     result={"path": str(tmp_path / "gone.csv"),
                             "filename": "gone.csv", "bytes": 4, "row_count": 1})
    processor, collection = build(tmp_path, [ready])
    processor.now = lambda: NOW

    result = await processor.reap()

    assert result["expired"] == 1
    assert result["files_deleted"] == 0
    assert collection.by_id("job-1")["status"] == jobs.STATUS_EXPIRED


@pytest.mark.asyncio
async def test_sweep_reaps_first_so_a_requeued_job_runs_on_the_same_tick(tmp_path):
    stale = make_job(status=jobs.STATUS_RUNNING, worker_id="worker-dead",
                     attempts=0, heartbeat_at=NOW - timedelta(minutes=10))
    processor, collection = build(tmp_path, [stale])
    processor.now = lambda: NOW

    result = await processor.sweep()

    assert result["reaped"]["requeued"] == 1
    assert result["claimed"] == ["job-1"]
    assert collection.by_id("job-1")["status"] == jobs.STATUS_READY


# --------------------------------------------------------------------------
# Registration - three places, or it runs once at boot and never again
# --------------------------------------------------------------------------

def _function_source(path, *names):
    """The source of one nested function, located by AST rather than by grep so a
    match in a neighbouring function cannot pass for a registration."""
    tree = ast.parse(path.read_text(encoding="utf-8"))
    node = tree
    for name in names:
        node = next(
            child for child in ast.walk(node)
            if isinstance(child, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef))
            and child.name == name
        )
    return ast.get_source_segment(path.read_text(encoding="utf-8"), node)


@pytest.mark.parametrize("path,names", [
    (CRON_MANAGER, ("CronManager", "_get_job_function")),
    (CRON_MANAGER, ("CronManager", "_register_fallback_jobs")),
    (JOB_SEEDER, ("SchedulerJobSeeder", "get_default_jobs")),
])
def test_the_sweep_job_is_registered_in_every_place_that_matters(path, names):
    """Missing any one of the three means the job runs once at boot and never
    fires again - the failure this whole comment block exists to prevent."""
    assert worker.JOB_TYPE in _function_source(path, *names)


def test_the_schedule_is_five_cron_fields():
    """_add_job_to_scheduler rejects anything else outright, and never passes
    second= to CronTrigger, so sub-minute is not expressible."""
    assert len(worker.JOB_SCHEDULE.split()) == 5
    assert worker.JOB_SCHEDULE == "* * * * *"


# --------------------------------------------------------------------------
# XLSX - the one format that cannot stream
# --------------------------------------------------------------------------

@pytest.mark.asyncio
async def test_an_xlsx_job_merges_every_chunk_into_one_workbook(tmp_path):
    """The sheet's columns are the dates, so a later chunk adds columns rather than
    rows: the parts have to be merged before the workbook is written, not appended
    to it one at a time."""
    from openpyxl import load_workbook

    processor, collection = build(
        tmp_path, [make_job(filters=make_filters(format="xlsx"))])

    result = await processor.run_job("job-1")

    assert result["status"] == jobs.STATUS_READY
    path = collection.by_id("job-1")["result"]["path"]
    assert path.endswith(".xlsx")

    book = load_workbook(path)
    assert "Filters" in book.sheetnames
    sheet = book[book.sheetnames[0]]
    header = [cell.value for cell in next(sheet.iter_rows(max_row=1))]
    assert header[:2] == ["Entity", "Block"]
    # One column per chunk date, all ten of them, in one sheet.
    assert len(header) == 2 + 10


@pytest.mark.asyncio
async def test_run_job_on_a_job_that_no_longer_exists_says_so(tmp_path):
    processor, _ = build(tmp_path, [])

    assert (await processor.run_job("ghost"))["status"] == worker.OUTCOME_MISSING
