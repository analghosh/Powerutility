"""The two export writers: streaming CSV and write-only XLSX.

Both write the same table (PTC, 2026-09-20): one row per (day, block), the identity
columns the chosen Group Date By calls for, then one column per series. The workbook
holds it on a single `Data` sheet; the CSV streams it.

Three properties are asserted here rather than assumed:

1. **The row stream stays lazy.** A long range is hundreds of thousands of rows and
   neither writer may hold them all. The single table does need every series before
   it can emit a row — they are its columns — so what is proved here is that the
   *rows* are produced one at a time, not that the series list is untouched.

2. **A missing value is an empty cell.** Never `0`, never the string `"None"`. The
   docx is explicit and so is the pivot: "nothing was scheduled for this block" and
   "zero was scheduled" are different facts, and a spreadsheet that renders the
   first as the second is a wrong number in front of a trader.

3. **The file says what it actually did.** The Filters block reports the grouping
   that ran, read from the response. It used to read the request, which is how a
   workbook of ISO-week columns came to carry the line `Date grouping: BLOCK_WISE`.
"""
import csv as csv_module
import inspect
import io
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from services.processing.app.v2.helpers import analytics_csv  # noqa: E402
from services.processing.app.v2.helpers import analytics_xlsx  # noqa: E402


FILTERS = {
    "datasets": ["ACTUAL_DEMAND", "DC_PLANT"],
    "from_date": "2026-04-01",
    "to_date": "2026-08-31",
    "from_block": 1,
    "to_block": 96,
    "date_grouping": "DATE_WISE",
    "block_grouping": "INDIVIDUAL",
    "plants": ["PLANT_A", "PLANT_B"],
    "exchanges": ["IEX"],
    "format": "csv",
}


def _series(dataset, label, values, entity=None, unit="MW", quality="complete"):
    return {"dataset": dataset, "label": label, "unit": unit, "entity": entity,
            "values": values, "quality": quality}


def _response(series, dates=("2026-04-01", "2026-04-02"), blocks=(1, 2)):
    return {
        "axes": {"dates": list(dates), "blocks": list(blocks)},
        "series": list(series),
        "meta": {"generated_at": "2026-09-17T06:00:00+00:00", "truncated": False,
                 "missing_datasets": [], "failed_datasets": []},
    }


def _lines(response, filters=None):
    return [line.rstrip("\n") for line in analytics_csv.csv_rows(response, filters or FILTERS)]


def _comments(lines):
    return [ln for ln in lines if ln.startswith("#")]


def _table(lines):
    body = [ln for ln in lines if not ln.startswith("#")]
    return list(csv_module.reader(io.StringIO("\n".join(body))))


class _CountingMatrix(list):
    """A `values` matrix that records which block rows were actually read."""

    def __init__(self, rows):
        super().__init__(rows)
        self.touched = []

    def __getitem__(self, index):
        self.touched.append(index)
        return super().__getitem__(index)


# ------------------------------------------------------------------ streaming

def test_csv_rows_is_a_generator():
    assert inspect.isgeneratorfunction(analytics_csv.csv_rows)
    rows = analytics_csv.csv_rows(_response([_series("ACTUAL_DEMAND", "Actual Demand",
                                                     [[1.0, 2.0], [3.0, 4.0]])]), FILTERS)
    assert inspect.isgenerator(rows)


def test_csv_rows_builds_one_row_at_a_time():
    """Rows are produced on demand, not materialised up front.

    The single-table layout needs every series to build a row — they are its
    columns — so the series list is read eagerly. The rows are not: pulling the
    header block and the first data line must touch only the first block of the
    first day, whatever the range holds.
    """
    matrix = _CountingMatrix([[1.0, 2.0], [3.0, 4.0]])
    response = _response([_series("ACTUAL_DEMAND", "Actual Demand", matrix)])

    rows = analytics_csv.csv_rows(response, FILTERS)
    assert inspect.isgenerator(rows)

    pulled = []
    for _ in range(50):
        pulled.append(next(rows))
        if not pulled[-1].startswith("#") and len(pulled) > 1 and pulled[-1][0].isdigit():
            break
    else:  # pragma: no cover - the loop finds a data row in well under 50 lines
        raise AssertionError(f"no data row was yielded lazily: {pulled}")

    assert set(matrix.touched) == {0}, (
        f"only block 1 should have been read for the first row, got {matrix.touched}")


def test_every_line_is_newline_terminated_so_it_can_be_streamed_as_is():
    response = _response([_series("ACTUAL_DEMAND", "Actual Demand", [[1.0, 2.0], [3.0, 4.0]])])
    assert all(line.endswith("\n") for line in analytics_csv.csv_rows(response, FILTERS))


# --------------------------------------------------------------- header block

def test_header_block_records_the_applied_filters():
    lines = _lines(_response([_series("ACTUAL_DEMAND", "Actual Demand",
                                      [[1.0, 2.0], [3.0, 4.0]])]))
    header = "\n".join(_comments(lines))

    assert "2026-04-01" in header and "2026-08-31" in header      # date range
    assert "1" in header and "96" in header                        # block range
    assert "DATE_WISE" in header                                   # date grouping
    assert "INDIVIDUAL" in header                                  # block grouping
    assert "ACTUAL_DEMAND" in header and "DC_PLANT" in header      # dataset list
    assert "PLANT_A" in header                                     # plant selection
    assert "2026-09-17T06:00:00+00:00" in header                   # generated_at


def test_header_block_records_custom_block_ranges_when_they_apply():
    filters = dict(FILTERS, block_grouping="CUSTOM", custom_blocks=[[1, 8], [9, 16]])
    header = "\n".join(_comments(_lines(_response([
        _series("ACTUAL_DEMAND", "Actual Demand", [[1.0, 2.0], [3.0, 4.0]])]), filters)))
    assert "CUSTOM" in header
    assert "1-8" in header and "9-16" in header


def test_column_header_follows_the_filter_block():
    response = _response([_series("ACTUAL_DEMAND", "Actual Demand",
                                  [[1.0, 2.0], [3.0, 4.0]])])
    table = _table(_lines(response))
    assert table[0] == ["Date", "Time block", "Actual Demand (MW)"]
    assert table[0] == analytics_csv.COLUMNS(response, FILTERS)


def test_the_grouping_chooses_the_identity_columns():
    """Each pill adds its own column and leaves the rows at one per (day, block)."""
    values = [[1.0, 2.0], [3.0, 4.0]]
    expected = {
        "BLOCK_DATE_WISE": ["Date", "Time block"],
        "HOURLY_WISE": ["Date", "Hourly", "Time block"],
        "WEEK_WISE": ["Week", "Date", "Time block"],
        "MONTH_WISE": ["Month", "Date", "Time block"],
        "YEAR_WISE": ["Year", "Date", "Time block"],
        # The two retired codes still arrive from stored rows.
        "BLOCK_WISE": ["Date", "Time block"],
        "DATE_WISE": ["Date", "Time block"],
    }
    for grouping, identity in expected.items():
        filters = dict(FILTERS, date_grouping=grouping)
        table = _table(_lines(_response([_series("ACTUAL_DEMAND", "Actual Demand", values)]),
                              filters))
        assert table[0][: len(identity)] == identity, grouping
        # Same number of rows whatever the pill: it groups nothing.
        assert len(table) == 1 + 4, grouping


def test_one_row_per_day_and_block_in_reading_order():
    values = [[1.0, 2.0], [3.0, 4.0]]
    table = _table(_lines(_response([_series("ACTUAL_DEMAND", "Actual Demand", values)])))
    assert [(r[0], r[1]) for r in table[1:]] == [
        ("2026-04-01", "1"), ("2026-04-01", "2"),
        ("2026-04-02", "1"), ("2026-04-02", "2"),
    ]


# ------------------------------------------------------------------ the cells

def test_a_missing_value_is_an_empty_cell_never_zero_and_never_the_string_none():
    #                block 1: [d1 missing, d2 real zero]
    values = [[None, 0.0], [2850.0, None]]
    lines = _lines(_response([_series("ACTUAL_DEMAND", "Actual Demand", values)]))
    table = _table(lines)
    head, body = table[0], table[1:]
    vi = head.index("Actual Demand (MW)")

    cells = {(r[0], r[1]): r[vi] for r in body}
    assert cells[("2026-04-01", "1")] == ""            # missing
    assert cells[("2026-04-02", "2")] == ""            # missing
    assert "None" not in [r[vi] for r in body]
    assert float(cells[("2026-04-02", "1")]) == 0.0    # a real zero survives
    assert float(cells[("2026-04-01", "2")]) == 2850.0


def test_each_entity_gets_its_own_column():
    """Two plants of one dataset must never share a heading.

    Side by side on the same row is the point of the single table, and that only
    works if the entity is part of the column's identity.
    """
    values = [[1.0, 2.0], [3.0, 4.0]]
    response = _response([
        _series("DC_PLANT", "Plant wise DC", values, entity="PLANT_A"),
        _series("DC_PLANT", "Plant wise DC", values, entity="PLANT_B"),
        _series("ACP_DAM", "DAM ACP", values, entity="IEX", unit="Rs/MWh"),
        _series("ACTUAL_DEMAND", "Actual Demand", values),
    ])
    head = _table(_lines(response))[0]
    assert head == [
        "Date", "Time block",
        "Plant wise DC · PLANT_A (MW)",
        "Plant wise DC · PLANT_B (MW)",
        "DAM ACP · IEX (Rs/MWh)",
        # a portfolio-wide series is just its label and unit — no "· None", no " ()"
        "Actual Demand (MW)",
    ]


def test_the_readable_entity_label_wins_over_the_machine_key():
    series = _series("WEEKLY_MAPE", "Weekly MAPE", [[1.0]], entity="DAF", unit="%")
    series["entity_label"] = "Day Ahead"
    assert analytics_csv.series_label(series) == "Weekly MAPE · Day Ahead (%)"


def test_a_heading_containing_a_comma_is_quoted_not_split():
    response = _response([_series("ACP_DAM", "DAM ACP, weighted", [[1.0, 2.0], [3.0, 4.0]],
                                  unit="Rs/MWh")])
    table = _table(_lines(response))
    assert table[0][2] == "DAM ACP, weighted (Rs/MWh)"
    assert all(len(r) == len(table[0]) for r in table)


def test_a_named_block_band_names_itself_in_the_time_block_column():
    """The user picked "Peak", so the column says Peak — not a code, not a range."""
    response = _response([_series("ACTUAL_DEMAND", "Actual Demand", [[1.0, 2.0]])],
                         blocks=("PEAK",))
    table = _table(_lines(response, dict(FILTERS, block_grouping="PEAK")))
    assert [r[1] for r in table[1:]] == ["Peak", "Peak"]


# ----------------------------------------------------------------------- xlsx

def _sheets(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def _at(row, index):
    """Cell value, tolerating openpyxl trimming trailing blanks off a stored row."""
    return row[index] if index < len(row) else None


def _rows_of(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        return [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    finally:
        wb.close()


def test_xlsx_has_one_data_sheet_plus_a_filters_sheet(tmp_path):
    """One table, not one sheet per dataset.

    The workbook used to put every dataset on its own sheet, which meant comparing
    two of them was a copy-paste job. They are columns on one sheet now, and the
    two plant entities are two of those columns rather than two rows.
    """
    values = [[1.0, 2.0], [3.0, 4.0]]
    response = _response([
        _series("ACTUAL_DEMAND", "Actual Demand", values),
        _series("DC_PLANT", "Plant wise DC", values, entity="PLANT_A"),
        _series("DC_PLANT", "Plant wise DC", values, entity="PLANT_B"),
    ])
    out = tmp_path / "export.xlsx"
    result = analytics_xlsx.write_xlsx(response, FILTERS, str(out))

    assert _sheets(str(out)) == [analytics_xlsx.DATA_SHEET, analytics_xlsx.FILTERS_SHEET]
    assert result["path"] == str(out)
    assert result["bytes"] == out.stat().st_size > 0
    assert result["row_count"] == 4      # 2 days x 2 blocks, whatever is ticked

    rows = _rows_of(str(out), analytics_xlsx.DATA_SHEET)
    assert rows[0] == ["Date", "Time block",
                       "Actual Demand (MW)",
                       "Plant wise DC · PLANT_A (MW)",
                       "Plant wise DC · PLANT_B (MW)"]
    assert [r[1] for r in rows[1:]] == [1, 2, 1, 2]


def test_the_date_column_holds_real_dates_not_text(tmp_path):
    """Text dates cannot be sorted, filtered or subtracted in Excel.

    Write-only mode has no styles to inherit, so the cells carry their own number
    format; without it the column renders as a serial number.
    """
    out = tmp_path / "export.xlsx"
    analytics_xlsx.write_xlsx(
        _response([_series("ACTUAL_DEMAND", "Actual Demand", [[1.0, 2.0], [3.0, 4.0]])]),
        FILTERS, str(out))

    wb = openpyxl.load_workbook(str(out))
    try:
        sheet = wb[analytics_xlsx.DATA_SHEET]
        cell = sheet.cell(row=2, column=1)
        assert isinstance(cell.value, datetime)
        assert cell.value.date() == date(2026, 4, 1)
        assert cell.number_format == analytics_xlsx.DATE_NUMBER_FORMAT
    finally:
        wb.close()


def test_the_grouping_column_is_written_into_the_sheet(tmp_path):
    out = tmp_path / "export.xlsx"
    analytics_xlsx.write_xlsx(
        _response([_series("ACTUAL_DEMAND", "Actual Demand", [[1.0, 2.0], [3.0, 4.0]])]),
        dict(FILTERS, date_grouping="WEEK_WISE"), str(out))

    rows = _rows_of(str(out), analytics_xlsx.DATA_SHEET)
    assert rows[0][:3] == ["Week", "Date", "Time block"]
    # 1 and 2 Apr 2026 are both ISO week 14.
    assert {r[0] for r in rows[1:]} == {"W14 2026"}


def test_two_datasets_with_the_same_long_label_still_get_distinct_columns(tmp_path):
    """Sheet-title truncation used to be the risk here; column headings are not
    capped at 31 characters, and the entity keeps two series of one dataset apart."""
    long_a = "Intra Day Forecast MAPE Northern Region A"
    long_b = "Intra Day Forecast MAPE Northern Region B"
    assert long_a[:31] == long_b[:31]
    values = [[1.0, 2.0], [3.0, 4.0]]
    out = tmp_path / "export.xlsx"
    analytics_xlsx.write_xlsx(
        _response([_series("MAPE_IDF_A", long_a, values),
                   _series("MAPE_IDF_B", long_b, values)]), FILTERS, str(out))

    head = _rows_of(str(out), analytics_xlsx.DATA_SHEET)[0]
    assert head[2:] == [f"{long_a} (MW)", f"{long_b} (MW)"]
    assert len(set(head)) == len(head)


def test_filters_sheet_records_the_applied_filters(tmp_path):
    out = tmp_path / "export.xlsx"
    analytics_xlsx.write_xlsx(
        _response([_series("ACTUAL_DEMAND", "Actual Demand", [[1.0, 2.0], [3.0, 4.0]])]),
        FILTERS, str(out))

    text = "\n".join(
        " ".join("" if c is None else str(c) for c in row)
        for row in _rows_of(str(out), analytics_xlsx.FILTERS_SHEET))
    for expected in ("2026-04-01", "2026-08-31", "96", "BLOCK_DATE_WISE", "INDIVIDUAL",
                     "ACTUAL_DEMAND", "2026-09-17T06:00:00+00:00"):
        assert expected in text, f"{expected!r} missing from Filters sheet:\n{text}"


def test_xlsx_missing_value_is_an_empty_cell_never_zero(tmp_path):
    #         block 1: [1 Apr missing, 2 Apr a real zero]
    values = [[None, 0.0], [2850.0, None]]
    out = tmp_path / "export.xlsx"
    analytics_xlsx.write_xlsx(
        _response([_series("ACTUAL_DEMAND", "Actual Demand", values)]), FILTERS, str(out))

    rows = _rows_of(str(out), analytics_xlsx.DATA_SHEET)
    head, body = rows[0], rows[1:]
    vi = head.index("Actual Demand (MW)")
    cells = {(r[0].date(), r[1]): _at(r, vi) for r in body}

    assert cells[(date(2026, 4, 1), 1)] is None      # missing, not 0 and not "None"
    assert cells[(date(2026, 4, 2), 1)] == 0.0       # a real zero survives
    assert cells[(date(2026, 4, 1), 2)] == 2850.0
    # A trailing blank is not stored in the sheet at all, which is the strongest
    # possible form of "empty": Excel shows no value, and no 0 was ever written.
    assert cells[(date(2026, 4, 2), 2)] is None
    assert all(c != "None" for row in body for c in row)


def test_write_xlsx_creates_a_missing_export_directory(tmp_path):
    out = tmp_path / "exports" / "nested" / "export.xlsx"
    assert not out.parent.exists()
    analytics_xlsx.write_xlsx(
        _response([_series("ACTUAL_DEMAND", "Actual Demand", [[1.0, 2.0], [3.0, 4.0]])]),
        FILTERS, str(out))
    assert out.exists()


def test_write_xlsx_reports_an_unwritable_directory_clearly(tmp_path):
    blocker = tmp_path / "exports"
    blocker.write_text("not a directory")           # makedirs cannot win here
    out = blocker / "export.xlsx"
    with pytest.raises(analytics_xlsx.ExportWriteError) as exc:
        analytics_xlsx.write_xlsx(
            _response([_series("ACTUAL_DEMAND", "Actual Demand", [[1.0, 2.0], [3.0, 4.0]])]),
            FILTERS, str(out))
    assert str(blocker) in str(exc.value)
    assert "ANALYTICS_EXPORT_DIR" in str(exc.value)


# ------------------------------------------------------------------- filename

def test_filename_for_embeds_the_module_name_and_the_date_range():
    name = analytics_csv.filename_for(FILTERS, "csv")
    assert name.endswith(".csv")
    assert "analytics" in name
    assert "2026-04-01" in name and "2026-08-31" in name
    assert " " not in name


def test_filename_for_uses_the_module_when_the_filters_name_one():
    name = analytics_csv.filename_for(dict(FILTERS, module="Demand Analytics"), "xlsx")
    assert name.endswith(".xlsx")
    assert "demand_analytics" in name


def test_a_user_supplied_name_replaces_the_generated_one():
    name = analytics_csv.filename_for(dict(FILTERS, file_name="Aug DC report"), "csv")
    assert name == "Aug_DC_report.csv"


def test_a_user_supplied_name_never_carries_its_own_extension():
    # `report.exe.csv` is a file a mail client may open on the first extension.
    assert analytics_csv.filename_for(
        dict(FILTERS, file_name="report.csv"), "csv") == "report.csv"
    assert analytics_csv.filename_for(
        dict(FILTERS, file_name="report.xlsx"), "csv") == "report.csv"


@pytest.mark.parametrize("supplied", [
    "../../etc/passwd",
    "..\\..\\windows\\system32\\evil",
    "sub/dir/report",
])
def test_a_user_supplied_name_cannot_escape_the_export_directory(supplied):
    import os
    name = analytics_csv.filename_for(dict(FILTERS, file_name=supplied), "csv")
    assert "/" not in name and "\\" not in name
    joined = os.path.normpath(os.path.join("/srv/exports", name))
    assert joined.startswith("/srv/exports" + os.sep)


@pytest.mark.parametrize("supplied", ['inject"; drop', "line\nbreak", "semi;colon"])
def test_a_user_supplied_name_cannot_inject_a_content_disposition_header(supplied):
    name = analytics_csv.filename_for(dict(FILTERS, file_name=supplied), "csv")
    assert not set(name) & set('";\r\n')


@pytest.mark.parametrize("supplied", ["", "   ", "....", "///"])
def test_an_unusable_name_falls_back_to_the_generated_one(supplied):
    # The file is what the user asked for; its name is not worth failing the export.
    name = analytics_csv.filename_for(dict(FILTERS, file_name=supplied), "csv")
    assert name == analytics_csv.filename_for(FILTERS, "csv")


def test_a_user_supplied_name_is_truncated_rather_than_rejected():
    name = analytics_csv.filename_for(dict(FILTERS, file_name="a" * 400), "csv")
    assert name == "a" * analytics_csv.MAX_FILENAME_STEM + ".csv"


def test_a_leading_dot_cannot_make_the_export_a_hidden_file():
    assert analytics_csv.filename_for(
        dict(FILTERS, file_name=".hidden"), "csv") == "hidden.csv"


def test_filename_for_rejects_an_unsupported_format():
    with pytest.raises(ValueError):
        analytics_csv.filename_for(FILTERS, "pdf")


def test_xlsx_reuses_the_same_filename_helper():
    assert analytics_xlsx.filename_for is analytics_csv.filename_for


def test_csv_header_names_partially_covered_periods():
    """Wed 26 → Fri 28 Aug rolled up weekly: the bucket says W35, the note says 3 of 7."""
    from services.processing.app.v2.helpers.analytics_csv import header_lines
    response = {"axes": {"dates": ["2026-W35"], "blocks": [1]}, "series": [],
                "meta": {"period_coverage": {"2026-W35": {
                    "days_covered": 3, "days_in_period": 7, "partial": True,
                    "from": "2026-08-26", "to": "2026-08-28"}}}}
    text = "".join(header_lines({"datasets": ["WEEKLY_FORECAST"]}, response))
    assert "# Partial periods: 2026-W35 (3 of 7 days)" in text


def test_csv_header_is_silent_when_every_period_is_complete():
    from services.processing.app.v2.helpers.analytics_csv import header_lines
    response = {"axes": {"dates": ["2026-W35"], "blocks": [1]}, "series": [],
                "meta": {"period_coverage": {"2026-W35": {
                    "days_covered": 7, "days_in_period": 7, "partial": False,
                    "from": "2026-08-24", "to": "2026-08-30"}}}}
    text = "".join(header_lines({"datasets": ["WEEKLY_FORECAST"]}, response))
    assert "Partial periods" not in text
