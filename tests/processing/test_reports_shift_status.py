"""
Tests for the Shift Status Report service, routes and download.

The figures are the dashboard's datasets plus the latest real-time forecast and the
shift desk's remarks; these tests pin how they are combined into EDP's sheet.
"""

import asyncio
import io
import uuid
from datetime import datetime, timedelta, timezone

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from unittest.mock import MagicMock

from libs.audit_queue.simple_logger import AuditLogger
from services.processing.app.core.dependencies import get_audit_logger
from services.processing.app.v2.api.router import reports_shift_status_routes as routes
from services.processing.app.v2.dependencies import get_user_context
from services.processing.app.v2.helpers import shift_status_xlsx
from services.processing.app.v2.services.reports_shift_status_service import (
    FULL_DAY_SHIFT_ID,
    RemarkLocked,
    ReportsShiftStatusService,
    ShiftNotFound,
    editable_dates,
    shift_person_dates,
)

SHIFTS = [
    {"id": "a", "shift_name": "Shift A", "shift_timing": "7-14"},
    {"id": "b", "shift_name": "Shift B", "shift_timing": "14-21"},
    {"id": "c", "shift_name": "Shift C", "shift_timing": "21-7"},
    {"id": "edp", "shift_name": "Evening", "shift_timing": "15:30-24:00"},
    {"id": "bad", "shift_name": "Broken", "shift_timing": "abc"},
]


class FakeAnalytics:
    """Answers `fetch` from `{(dataset, entity): {(iso, block): value}}`."""

    def __init__(self, grids):
        self.grids = grids
        self.requests = []

    async def fetch(self, request):
        self.requests.append(request)
        day, last = request.from_date, request.to_date  # dates, parsed by the model
        dates = []
        while day <= last:
            dates.append(day.isoformat())
            day += timedelta(days=1)
        blocks = list(range(1, 97))
        series = []
        for code in request.datasets:
            entities = request.plants if code == "SG_PLANT" else [None]
            for entity in entities:
                grid = self.grids.get((code, entity), {})
                series.append({
                    "dataset": code,
                    "entity": entity,
                    "values": [[grid.get((d, b)) for d in dates] for b in blocks],
                })
        return {"axes": {"dates": dates, "blocks": blocks}, "series": series,
                "meta": {"failed_datasets": []}}


class FakeForecast:
    def __init__(self, docs=()):
        self.docs = list(docs)
        self.calls = []

    async def find_by_date_range_and_block_range(self, from_date, to_date, from_block,
                                                 to_block, market_type):
        self.calls.append((from_date, to_date, from_block, to_block, market_type))
        return self.docs


class FakeEntries:
    def __init__(self, stored=None):
        self.stored = dict(stored or {})
        self.saves = []
        self.person_saves = []

    async def find_for_dates(self, dates):
        wanted = {d.isoformat() for d in dates}
        return {k: v for k, v in self.stored.items() if k[0] in wanted}

    async def insert_new(self, entries, *, user_id, username):
        self.saves.append((entries, user_id, username))
        locked = [{"date": e["date"], "time_block": e["time_block"]} for e in entries
                  if (e["date"], e["time_block"]) in self.stored]
        for e in entries:
            self.stored.setdefault((e["date"], e["time_block"]), {
                "remark": e["remark"], "updated_by_name": username,
                "updated_at": datetime(2026, 9, 28, 18, 0)})
        return {"saved": len(entries) - len(locked), "locked": locked}

    async def set_shift_person(self, iso_date, names, *, user_id, username):
        self.person_saves.append((iso_date, names, user_id, username))
        doc = {"kind": "shift_person", "names": names, "updated_by_name": username,
               "updated_at": datetime(2026, 9, 29, 9, 0, tzinfo=timezone.utc)}
        self.stored[(iso_date, 0)] = doc
        return {"date": iso_date, **doc}


def _loader(shifts=SHIFTS):
    async def load():
        return [dict(s) for s in shifts]
    return load


# 28 Sep 2026 block 63 and 69 on the test site, and EDP's sheet for the same blocks.
GRIDS = {
    ("FORECAST_DAF", None): {("2026-09-28", 63): 451.3, ("2026-09-28", 69): 457.8},
    ("FORECAST_IDF", None): {("2026-09-28", 63): 442.5, ("2026-09-28", 69): 461.7},
    ("SG_TOTAL", None): {("2026-09-28", 63): 482.45, ("2026-09-28", 69): 495.25,
                         ("2026-09-29", 1): 500.0},
    ("SG_PLANT", "PPCLGPP"): {("2026-09-28", 63): 15.0, ("2026-09-28", 69): 15.0},
    ("DRAWAL", None): {("2026-09-28", 69): 480.0, ("2026-09-28", 70): 0.0},
    ("PF_BIDDED_DAM", None): {("2026-09-28", 63): -60.0},
    ("PF_TRADED_DAM", None): {("2026-09-28", 63): -51.3},
    ("PF_BIDDED_RTM", None): {("2026-09-28", 67): 40.0},
    ("PF_TRADED_RTM", None): {("2026-09-28", 67): 23.93},
}

FORECAST_DOCS = [
    {"scheduled_at": datetime(2026, 9, 28), "time_block": 63, "forecast": 468.5,
     "source": "rtm_forecast_api"},
    {"scheduled_at": datetime(2026, 9, 28), "time_block": 64, "forecast": 466.3,
     "source": "rtm_forecast_api"},
    {"scheduled_at": datetime(2026, 9, 28), "time_block": 64, "forecast": 470.0,
     "source": "csv"},
]


TODAY = datetime(2026, 9, 29).date()  # remarks open for 27, 28, 29 and 30 Sep


def _service(grids=GRIDS, docs=FORECAST_DOCS, remarks=None, shifts=SHIFTS, today=TODAY):
    return ReportsShiftStatusService(
        analytics=FakeAnalytics(grids), forecast_repository=FakeForecast(docs),
        entries_repository=FakeEntries(remarks), shifts_loader=_loader(shifts),
        today=lambda: today)


def _report(service=None, from_date="2026-09-28", to_date="2026-09-28", shift_id="edp"):
    return asyncio.run((service or _service()).get_report(from_date, to_date, shift_id))


def _block(report, iso, block):
    for section in report["sections"]:
        for row in section["blocks"]:
            if row["date"] == iso and row["time_block"] == block:
                return row
    raise AssertionError(f"no row {iso} #{block}")


# ------------------------------------------------------------------ shifts

def test_the_picker_lists_configured_shifts_then_full_day_and_skips_a_broken_one():
    options = asyncio.run(_service().shifts())
    assert [o["id"] for o in options] == ["a", "b", "c", "edp", FULL_DAY_SHIFT_ID]


def test_an_unknown_shift_is_refused():
    with pytest.raises(ShiftNotFound):
        _report(shift_id="nope")


def test_the_edp_evening_shift_is_one_section_of_blocks_63_to_96():
    report = _report()
    assert len(report["sections"]) == 1
    blocks = report["sections"][0]["blocks"]
    assert [b["time_block"] for b in blocks] == list(range(63, 97))
    assert (blocks[0]["time_from"], blocks[0]["time_to"]) == ("15:30", "15:45")
    assert blocks[-1]["time_to"] == "00:00"


def test_a_night_shift_section_carries_the_next_days_blocks_with_their_own_date():
    service = _service()
    report = _report(service, shift_id="c")
    rows = report["sections"][0]["blocks"]
    assert report["sections"][0]["date"] == "2026-09-28"
    assert (rows[0]["date"], rows[0]["time_block"]) == ("2026-09-28", 85)
    assert (rows[-1]["date"], rows[-1]["time_block"]) == ("2026-09-29", 28)
    # The figures for the 29th are read too.
    assert _block(report, "2026-09-29", 1)["net_schedule"] == 500.0
    assert str(service.analytics.requests[0].to_date) == "2026-09-29"


def test_a_range_gives_one_section_per_date_in_order():
    report = _report(from_date="2026-09-26", to_date="2026-09-28", shift_id="b")
    assert [s["date"] for s in report["sections"]] == [
        "2026-09-26", "2026-09-27", "2026-09-28"]


def test_full_day_is_all_96_blocks():
    report = _report(shift_id=FULL_DAY_SHIFT_ID)
    assert len(report["sections"][0]["blocks"]) == 96


def test_a_range_over_31_days_is_refused():
    with pytest.raises(ValueError):
        _report(from_date="2026-08-01", to_date="2026-09-28")


# ----------------------------------------------------------------- figures

def test_net_schedule_is_final_schedule_less_the_intra_state_plant():
    row = _block(_report(), "2026-09-28", 69)
    assert row["net_schedule"] == 480.25  # EDP's sheet: 480


def test_the_intra_state_plant_is_read_in_its_own_query():
    # The dashboard's plant filter also narrows Final Schedule, so the two
    # must never share a request.
    service = _service()
    _report(service)
    main, plants = service.analytics.requests
    assert "SG_TOTAL" in main.datasets and not main.plants
    assert plants.datasets == ["SG_PLANT"] and plants.plants == ["PPCLGPP"]


def test_drawal_and_deviation():
    row = _block(_report(), "2026-09-28", 69)
    assert row["actual_drawal"] == 480.0
    assert row["deviation"] == pytest.approx(-0.25)


def test_a_zero_drawal_is_blank_and_so_is_its_deviation():
    row = _block(_report(), "2026-09-28", 70)
    assert row["actual_drawal"] is None
    assert row["deviation"] is None
    assert row["needs_remark"] is False


def test_the_real_time_forecast_prefers_ptcs_upload_then_the_api_then_intra_day():
    report = _report()
    assert (_block(report, "2026-09-28", 64)["real_time_forecast"],
            _block(report, "2026-09-28", 64)["real_time_forecast_source"]) == (470.0, "Manual")
    assert (_block(report, "2026-09-28", 63)["real_time_forecast"],
            _block(report, "2026-09-28", 63)["real_time_forecast_source"]) == (468.5, "Auto")
    assert (_block(report, "2026-09-28", 69)["real_time_forecast"],
            _block(report, "2026-09-28", 69)["real_time_forecast_source"]) == (461.7, "Intra Day")


def test_bid_and_traded_are_both_shown():
    report = _report()
    assert (_block(report, "2026-09-28", 63)["idam_bid"],
            _block(report, "2026-09-28", 63)["idam_traded"]) == (-60.0, -51.3)
    assert (_block(report, "2026-09-28", 67)["rtm_bid"],
            _block(report, "2026-09-28", 67)["rtm_traded"]) == (40.0, 23.93)
    assert _block(report, "2026-09-28", 70)["rtm_bid"] is None


def test_idam_adds_green_and_high_price_trades():
    grids = dict(GRIDS)
    grids[("PF_TRADED_GDAM", None)] = {("2026-09-28", 63): -10.0}
    assert _block(_report(_service(grids)), "2026-09-28", 63)["idam_traded"] == -61.3


def test_a_day_where_gdam_and_hpdam_copy_dam_counts_once_and_says_so():
    grids = dict(GRIDS)
    for market in ("GDAM", "HPDAM"):
        grids[(f"PF_BIDDED_{market}", None)] = dict(GRIDS[("PF_BIDDED_DAM", None)])
        grids[(f"PF_TRADED_{market}", None)] = dict(GRIDS[("PF_TRADED_DAM", None)])
    report = _report(_service(grids))
    assert _block(report, "2026-09-28", 63)["idam_bid"] == -60.0
    assert any("GDAM, HPDAM on 2026-09-28" in w for w in report["warnings"])


# ----------------------------------------------------------------- remarks

def _remark(text, name, minute):
    return {"remark": text, "updated_by_name": name,
            "updated_at": datetime(2026, 9, 28, 18, minute, tzinfo=timezone.utc)}


def test_a_block_beyond_15_mw_without_a_remark_is_flagged_and_counted():
    grids = dict(GRIDS)
    grids[("DRAWAL", None)] = {("2026-09-28", 63): 484.0, ("2026-09-28", 69): 511.0}
    remarks = {("2026-09-28", 69): _remark("Load increased", "sugumar", 5)}
    report = _report(_service(grids, remarks=remarks))
    assert _block(report, "2026-09-28", 63)["needs_remark"] is True   # +16.55
    assert _block(report, "2026-09-28", 69)["needs_remark"] is False  # remark given
    section = report["sections"][0]
    assert section["over_threshold_count"] == 2
    assert section["needs_remark_count"] == 1
    assert section["drawal_block_count"] == 2


def test_shift_person_is_whoever_saved_the_latest_remark_in_the_section():
    remarks = {
        ("2026-09-28", 63): _remark("a", "sugumar", 5),
        ("2026-09-28", 70): _remark("b", "ravi", 40),
        ("2026-09-28", 10): _remark("outside the shift", "later", 59),
    }
    report = _report(_service(remarks=remarks))
    assert report["sections"][0]["shift_person"] == "ravi"
    assert _block(report, "2026-09-28", 63)["remark"] == "a"


def test_no_remark_means_no_shift_person():
    section = _report()["sections"][0]
    assert section["shift_person"] is None
    assert section["shift_person_source"] is None


# ------------------------------------------------------------ shift person

def _names(names, by="sugumar"):
    return {"kind": "shift_person", "names": names, "updated_by_name": by,
            "updated_at": datetime(2026, 9, 28, 16, 0, tzinfo=timezone.utc)}


def test_typed_names_win_over_the_latest_remark_writer():
    stored = {("2026-09-28", 0): _names("ajay, vijay"),
              ("2026-09-28", 70): _remark("b", "ravi", 40)}
    section = _report(_service(remarks=stored))["sections"][0]
    assert section["shift_person"] == "ajay, vijay"
    assert section["shift_person_source"] == "entered"
    assert section["shift_person_updated_by"] == "sugumar"
    assert section["shift_person_updated_at"] == "2026-09-28T16:00:00+00:00"
    assert section["remark_count"] == 1  # the names are not a remark


def test_without_typed_names_the_latest_remark_writer_is_the_fallback():
    stored = {("2026-09-28", 70): _remark("b", "ravi", 40)}
    section = _report(_service(remarks=stored))["sections"][0]
    assert section["shift_person"] == "ravi"
    assert section["shift_person_source"] == "remarks"
    assert section["shift_person_updated_by"] is None


@pytest.mark.parametrize("shift_id", ["a", "c", "edp", FULL_DAY_SHIFT_ID])
def test_one_set_of_names_per_date_shows_in_every_shift(shift_id):
    stored = {("2026-09-28", 0): _names("ajay, vijay")}
    report = _report(_service(remarks=stored), shift_id=shift_id)
    assert report["sections"][0]["shift_person"] == "ajay, vijay"


def test_names_are_read_in_the_same_query_as_the_remarks():
    calls = []

    class Counting(FakeEntries):
        async def find_for_dates(self, dates):
            calls.append(sorted(d.isoformat() for d in dates))
            return await super().find_for_dates(dates)

    service = _service()
    service._entries_repository = Counting({("2026-09-28", 0): _names("ajay")})
    report = asyncio.run(service.get_report("2026-09-28", "2026-09-28", "c"))
    assert report["sections"][0]["shift_person"] == "ajay"
    assert calls == [["2026-09-28", "2026-09-29"]]


def test_the_shift_person_window_is_today_and_tomorrow():
    assert shift_person_dates(TODAY) == {"2026-09-29", "2026-09-30"}
    report = _report(_service(), from_date="2026-09-27", to_date="2026-09-30")
    assert [s["shift_person_editable"] for s in report["sections"]] == [
        False, False, True, True]
    assert report["shift_person_dates"] == ["2026-09-29", "2026-09-30"]


@pytest.mark.parametrize("day, names, message", [
    ("29-09-2026", "ajay", "YYYY-MM-DD"),
    ("2026-09-29", 5, "text"),
    ("2026-09-29", None, "text"),
    ("2026-09-29", " , ", "at least one name"),
    ("2026-09-29", "x" * 201, "200"),
    ("2026-09-28", "ajay", "today and tomorrow"),
    ("2026-10-01", "ajay", "today and tomorrow"),
])
def test_names_the_screen_could_not_have_sent_are_refused(day, names, message):
    with pytest.raises(ValueError, match=message):
        asyncio.run(_service().save_shift_person(day, names, user_id="u1",
                                                 username="sugumar"))


def test_names_are_tidied_and_can_be_changed_again_and_again():
    service = _service()
    first = asyncio.run(service.save_shift_person(
        "2026-09-29", "ajay,vijay ", user_id="u1", username="sugumar"))
    assert first["shift_person"] == "ajay, vijay"
    asyncio.run(service.save_shift_person(
        "2026-09-29", "ajay, vijay, ravi", user_id="u2", username="ravi"))
    saves = service.entries_repository.person_saves
    assert saves == [("2026-09-29", "ajay, vijay", "u1", "sugumar"),
                     ("2026-09-29", "ajay, vijay, ravi", "u2", "ravi")]
    report = _report(service, from_date="2026-09-29", to_date="2026-09-29")
    assert report["sections"][0]["shift_person"] == "ajay, vijay, ravi"


def test_a_date_with_names_still_takes_remarks():
    service = _service(remarks={("2026-09-28", 0): _names("ajay")})
    result = asyncio.run(service.save_remarks(
        [{"date": "2026-09-28", "time_block": 67, "remark": "Buy bid not cleared"}],
        user_id="u1", username="sugumar"))
    assert result == {"saved": 1, "blocks": 1}


OPEN = editable_dates(TODAY)


@pytest.mark.parametrize("entries, message", [
    ([], "non-empty"),
    ("x", "non-empty"),
    ([{"date": "28-09-2026", "time_block": 1, "remark": "x"}], "YYYY-MM-DD"),
    ([{"date": "2026-09-28", "time_block": 97, "remark": "x"}], "1 to 96"),
    ([{"date": "2026-09-28", "time_block": True, "remark": "x"}], "1 to 96"),
    ([{"date": "2026-09-28", "time_block": 1, "remark": 5}], "text"),
    ([{"date": "2026-09-28", "time_block": 1, "remark": "   "}], "cannot be empty"),
    ([{"date": "2026-09-28", "time_block": 1, "remark": "x" * 501}], "500"),
    ([{"date": "2026-09-28", "time_block": 1, "remark": "a"},
      {"date": "2026-09-28", "time_block": 1, "remark": "b"}], "twice"),
])
def test_remarks_the_screen_could_not_have_sent_are_refused(entries, message):
    with pytest.raises(ValueError, match=message):
        ReportsShiftStatusService.validate_remarks(entries, OPEN)


def test_the_remark_window_is_tomorrow_today_yesterday_and_the_day_before():
    assert editable_dates(TODAY) == {"2026-09-27", "2026-09-28", "2026-09-29", "2026-09-30"}


@pytest.mark.parametrize("day", ["2026-09-26", "2026-10-01"])
def test_a_remark_outside_the_window_is_refused(day):
    with pytest.raises(ValueError, match="tomorrow, today, yesterday and the day before"):
        ReportsShiftStatusService.validate_remarks(
            [{"date": day, "time_block": 1, "remark": "x"}], OPEN)


def test_rows_say_whether_they_can_still_take_a_remark():
    remarks = {("2026-09-28", 63): _remark("done", "sugumar", 5)}
    report = _report(_service(remarks=remarks))
    assert _block(report, "2026-09-28", 63)["remark_editable"] is False  # submitted
    assert _block(report, "2026-09-28", 64)["remark_editable"] is True
    assert report["remark_dates"] == ["2026-09-27", "2026-09-28", "2026-09-29", "2026-09-30"]
    old = _report(_service(today=datetime(2026, 10, 5).date()))
    assert old["sections"][0]["editable_count"] == 0


def test_a_night_shift_row_on_the_next_day_uses_its_own_date_for_the_window():
    report = _report(_service(today=datetime(2026, 10, 1).date()), shift_id="c")
    assert _block(report, "2026-09-28", 90)["remark_editable"] is False  # 3 days back
    assert _block(report, "2026-09-29", 1)["remark_editable"] is True


def test_tomorrows_blocks_already_take_a_remark():
    report = _report(_service(), from_date="2026-09-30", to_date="2026-09-30")
    assert report["sections"][0]["editable_count"] == 34
    result = asyncio.run(_service().save_remarks(
        [{"date": "2026-09-30", "time_block": 63, "remark": "Planned outage"}],
        user_id="u1", username="sugumar"))
    assert result == {"saved": 1, "blocks": 1}


def test_submitting_records_who():
    service = _service()
    result = asyncio.run(service.save_remarks(
        [{"date": "2026-09-28", "time_block": 67, "remark": " Buy bid not cleared "}],
        user_id="u1", username="sugumar"))
    assert result == {"saved": 1, "blocks": 1}
    entries, user_id, username = service.entries_repository.saves[0]
    assert (user_id, username) == ("u1", "sugumar")
    assert entries[0]["remark"] == "Buy bid not cleared"


def test_a_submitted_remark_cannot_be_changed_and_nothing_else_is_written():
    service = _service(remarks={("2026-09-28", 67): _remark("first", "sugumar", 5)})
    with pytest.raises(RemarkLocked) as caught:
        asyncio.run(service.save_remarks(
            [{"date": "2026-09-28", "time_block": 67, "remark": "second"},
             {"date": "2026-09-28", "time_block": 68, "remark": "new"}],
            user_id="u2", username="ravi"))
    assert caught.value.blocks == [{"date": "2026-09-28", "time_block": 67}]
    assert service.entries_repository.saves == []
    assert service.entries_repository.stored[("2026-09-28", 67)]["remark"] == "first"


# ---------------------------------------------------------------- download

def _sheet(report):
    return load_workbook(io.BytesIO(shift_status_xlsx.build_workbook(report))).active


def test_the_download_follows_the_shift_sheet_layout():
    remarks = {("2026-09-28", 63): _remark("Buy bid not cleared completely", "sugumar", 5)}
    ws = _sheet(_report(_service(remarks=remarks)))
    assert ws["A1"].value.startswith("Shift Status Report — 28-09-2026 · Evening")
    assert ws["A2"].value == "Shift Person Name" and ws["E2"].value == "sugumar"
    assert [ws.cell(row=3, column=c).value for c in (1, 2, 4, 5, 10, 12)] == [
        "Date", "Time", "Block", "DAY Ahead Forecast", "I-DAM Status", "RTM Status"]
    assert [ws.cell(row=4, column=c).value for c in (2, 3, 10, 11, 12, 13)] == [
        "From", "To", "Bid", "Traded", "Bid", "Traded"]
    merged = {str(r) for r in ws.merged_cells.ranges}
    assert {"A3:A4", "B3:C3", "J3:K3", "L3:M3", "N3:N4", "A2:D2"} <= merged
    first = [ws.cell(row=5, column=c).value for c in range(1, 15)]
    assert first[:4] == ["28-09-2026", "15:30", "15:45", 63]
    assert first[9:11] == [-60.0, -51.3]
    assert first[13] == "Buy bid not cleared completely"
    assert ws.cell(row=5 + 33, column=4).value == 96


def test_the_download_carries_the_typed_names():
    stored = {("2026-09-28", 0): _names("ajay, vijay"),
              ("2026-09-28", 63): _remark("x", "sugumar", 5)}
    ws = _sheet(_report(_service(remarks=stored)))
    assert ws["E2"].value == "ajay, vijay"


def test_each_date_gets_its_own_section_with_a_gap_between():
    ws = _sheet(_report(from_date="2026-09-27", to_date="2026-09-28"))
    # 1 title + 1 person + 2 header + 34 blocks + 1 blank = 39 rows per section.
    assert ws["A40"].value.startswith("Shift Status Report — 28-09-2026")
    assert ws.cell(row=39, column=1).value is None


def test_a_deviation_beyond_15_mw_is_shaded():
    grids = dict(GRIDS)
    grids[("DRAWAL", None)] = {("2026-09-28", 63): 484.0}
    ws = _sheet(_report(_service(grids)))
    assert ws.cell(row=5, column=9).fill.fgColor.rgb.endswith("FCE4D6")
    assert ws.cell(row=6, column=9).fill.fill_type is None


def test_the_filename_names_the_range_and_shift():
    assert shift_status_xlsx.filename_for("2026-09-28", "2026-09-28", "Shift B") == \
        "Shift-Status-Report-2026-09-28-Shift-B.xlsx"
    assert shift_status_xlsx.filename_for("2026-09-01", "2026-09-28", "Full day") == \
        "Shift-Status-Report-2026-09-01-to-2026-09-28-Full-day.xlsx"


# ------------------------------------------------------------------ routes

@pytest.fixture
def client(monkeypatch):
    service = _service()
    monkeypatch.setattr(routes, "reports_shift_status_service", service)
    app = FastAPI()
    app.include_router(routes.router)
    app.dependency_overrides[get_audit_logger] = lambda: MagicMock(spec=AuditLogger)
    app.dependency_overrides[get_user_context] = lambda: {
        "user_id": "u1", "username": "sugumar", "correlation_id": "t"}
    test_client = TestClient(app)
    test_client.service = service
    return test_client


def test_the_report_route_returns_sections(client):
    response = client.get("/reports-shift-status/", params={
        "from_date": "2026-09-28", "to_date": "2026-09-28", "shift_id": "edp"})
    assert response.status_code == 200
    assert response.json()["sections"][0]["block_count"] == 34


def test_an_unknown_shift_is_404_and_a_long_range_422(client):
    base = {"from_date": "2026-09-28", "to_date": "2026-09-28"}
    assert client.get("/reports-shift-status/", params={**base, "shift_id": "x"}).status_code == 404
    assert client.get("/reports-shift-status/", params={
        "from_date": "2026-08-01", "to_date": "2026-09-28"}).status_code == 422


def test_the_download_route_serves_an_xlsx(client):
    response = client.get("/reports-shift-status/download", params={
        "from_date": "2026-09-28", "to_date": "2026-09-28", "shift_id": "b"})
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(routes.XLSX_MEDIA_TYPE)
    assert "Shift-Status-Report-2026-09-28-Shift-B.xlsx" in \
        response.headers["content-disposition"]
    assert load_workbook(io.BytesIO(response.content)).active["A1"].value


def test_the_remarks_route_saves_as_the_signed_in_user(client):
    response = client.put("/reports-shift-status/remarks", json={"remarks": [
        {"date": "2026-09-28", "time_block": 67, "remark": "Buy bid not cleared"}]})
    assert response.status_code == 200
    assert client.service.entries_repository.saves[0][2] == "sugumar"


def test_the_remarks_route_answers_409_for_a_submitted_block(client):
    body = {"remarks": [{"date": "2026-09-28", "time_block": 70, "remark": "x"}]}
    assert client.put("/reports-shift-status/remarks", json=body).status_code == 200
    again = client.put("/reports-shift-status/remarks", json=body)
    assert again.status_code == 409
    assert again.json()["detail"]["locked"] == [{"date": "2026-09-28", "time_block": 70}]


def test_the_remarks_route_refuses_a_date_outside_the_window(client):
    response = client.put("/reports-shift-status/remarks", json={"remarks": [
        {"date": "2026-09-20", "time_block": 1, "remark": "x"}]})
    assert response.status_code == 422


def test_the_remarks_route_refuses_a_bad_block(client):
    response = client.put("/reports-shift-status/remarks", json={"remarks": [
        {"date": "2026-09-28", "time_block": 0, "remark": "x"}]})
    assert response.status_code == 422


def test_an_anonymous_caller_cannot_save_remarks(client):
    client.app.dependency_overrides[get_user_context] = lambda: {
        "user_id": "anonymous", "username": "anonymous"}
    response = client.put("/reports-shift-status/remarks", json={"remarks": [
        {"date": "2026-09-28", "time_block": 1, "remark": "x"}]})
    assert response.status_code == 401


def test_the_shift_person_route_saves_as_the_signed_in_user(client):
    response = client.put("/reports-shift-status/shift-person", json={
        "date": "2026-09-29", "names": "ajay,vijay"})
    assert response.status_code == 200
    assert response.json()["shift_person"] == "ajay, vijay"
    assert client.service.entries_repository.person_saves[0][3] == "sugumar"


@pytest.mark.parametrize("body", [
    {"date": "2026-09-27", "names": "ajay"},
    {"date": "2026-09-29", "names": "  "},
    {"date": "2026-09-29"},
])
def test_the_shift_person_route_refuses_bad_input(client, body):
    assert client.put("/reports-shift-status/shift-person", json=body).status_code == 422
    assert client.service.entries_repository.person_saves == []


def test_an_anonymous_caller_cannot_save_names(client):
    client.app.dependency_overrides[get_user_context] = lambda: {
        "user_id": "anonymous", "username": "anonymous"}
    response = client.put("/reports-shift-status/shift-person", json={
        "date": "2026-09-29", "names": "ajay"})
    assert response.status_code == 401


# ------------------------------------------------------- remarks on real Mongo

def test_remarks_are_written_once_on_real_mongo():
    from motor.motor_asyncio import AsyncIOMotorClient
    from services.processing.app.core.config import settings
    from services.processing.app.v2.repository.shift_status_entries_repository import (
        ShiftStatusEntriesRepository,
    )
    from datetime import date

    async def run():
        mongo = AsyncIOMotorClient(settings.MONGODB_URI, serverSelectionTimeoutMS=1500)
        try:
            await mongo.admin.command("ping")
        except Exception:
            pytest.skip("MongoDB not reachable")
        collection = mongo[settings.MONGODB_DATABASE][f"zz_shift_remarks_{uuid.uuid4().hex}"]
        repo = ShiftStatusEntriesRepository(legacy_collection_name=None)
        repo.collection = collection
        try:
            first = await repo.insert_new(
                [{"date": "2026-09-28", "time_block": 67, "remark": " first "}],
                user_id="u1", username="sugumar")
            assert first == {"saved": 1, "locked": []}
            second = await repo.insert_new(
                [{"date": "2026-09-28", "time_block": 67, "remark": "second"},
                 {"date": "2026-09-28", "time_block": 68, "remark": "x"}],
                user_id="u2", username="ravi")
            assert second == {"saved": 1,
                              "locked": [{"date": "2026-09-28", "time_block": 67}]}
            found = await repo.find_for_dates([date(2026, 9, 28)])
            assert found[("2026-09-28", 67)]["remark"] == "first"
            assert found[("2026-09-28", 67)]["updated_by_name"] == "sugumar"
            assert found[("2026-09-28", 68)]["updated_by_name"] == "ravi"
            assert await collection.count_documents({}) == 2
            indexes = await collection.index_information()
            assert indexes["uniq_date_time_block"]["unique"] is True
        finally:
            await collection.drop()
            mongo.close()

    asyncio.run(run())


async def _mongo_or_skip():
    from motor.motor_asyncio import AsyncIOMotorClient
    from services.processing.app.core.config import settings

    mongo = AsyncIOMotorClient(settings.MONGODB_URI, serverSelectionTimeoutMS=1500)
    try:
        await mongo.admin.command("ping")
    except Exception:
        pytest.skip("MongoDB not reachable")
    return mongo, mongo[settings.MONGODB_DATABASE]


def test_names_are_overwritten_and_sit_beside_the_remarks_on_real_mongo():
    from services.processing.app.v2.repository.shift_status_entries_repository import (
        ShiftStatusEntriesRepository,
    )
    from datetime import date

    async def run():
        mongo, database = await _mongo_or_skip()
        collection = database[f"zz_shift_entries_{uuid.uuid4().hex}"]
        repo = ShiftStatusEntriesRepository(legacy_collection_name=None)
        repo.collection = collection
        try:
            await repo.insert_new(
                [{"date": "2026-09-29", "time_block": 1, "remark": "x"}],
                user_id="u1", username="sugumar")
            await repo.set_shift_person("2026-09-29", "ajay", user_id="u1",
                                        username="sugumar")
            await repo.set_shift_person("2026-09-29", "ajay, vijay", user_id="u2",
                                        username="ravi")
            found = await repo.find_for_dates([date(2026, 9, 29)])
            assert found[("2026-09-29", 0)]["names"] == "ajay, vijay"
            assert found[("2026-09-29", 0)]["updated_by_name"] == "ravi"
            assert found[("2026-09-29", 0)]["kind"] == "shift_person"
            assert found[("2026-09-29", 1)]["kind"] == "remark"
            assert await collection.count_documents({}) == 2
        finally:
            await collection.drop()
            mongo.close()

    asyncio.run(run())


def test_the_old_remarks_collection_is_renamed_on_first_use_on_real_mongo():
    from services.processing.app.v2.repository.shift_status_entries_repository import (
        ShiftStatusEntriesRepository,
    )
    from datetime import date

    async def run():
        mongo, database = await _mongo_or_skip()
        suffix = uuid.uuid4().hex
        legacy = database[f"zz_shift_remarks_{suffix}"]
        target = database[f"zz_shift_entries_{suffix}"]
        await legacy.insert_one({"date": "2026-09-28", "time_block": 67,
                                 "remark": "kept", "updated_by_name": "sugumar"})
        repo = ShiftStatusEntriesRepository(legacy_collection_name=legacy.name)
        repo.collection = target
        try:
            found = await repo.find_for_dates([date(2026, 9, 28)])
            assert found[("2026-09-28", 67)]["remark"] == "kept"
            names = await database.list_collection_names(
                filter={"name": {"$in": [legacy.name, target.name]}})
            assert names == [target.name]
            assert (await target.index_information())["uniq_date_time_block"]["unique"]
            # A second process finds nothing left to rename.
            again = ShiftStatusEntriesRepository(legacy_collection_name=legacy.name)
            again.collection = target
            assert (await again.find_for_dates([date(2026, 9, 28)])).keys() == found.keys()
        finally:
            await legacy.drop()
            await target.drop()
            mongo.close()

    asyncio.run(run())


def test_when_a_remark_was_saved_is_sent_with_its_time_zone():
    naive_utc = {"remark": "x", "updated_by_name": "sugumar",
                 "updated_at": datetime(2026, 9, 28, 12, 0)}
    row = _block(_report(_service(remarks={("2026-09-28", 63): naive_utc})), "2026-09-28", 63)
    assert row["remark_updated_at"] == "2026-09-28T12:00:00+00:00"
