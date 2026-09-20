"""Write-only XLSX for the analytics dashboard export.

`openpyxl`'s **write-only mode** is not an optimisation here, it is the reason this
tier exists at all: a normal `Workbook` keeps every `Cell` object alive until `save()`,
which for 150k cells is hundreds of megabytes on a box that also serves bids. In
write-only mode each appended row is streamed to a temporary file and dropped, so the
peak cost is one row.

    wb = Workbook(write_only=True)          # streams rows to a temp file

The layout is one **Data** sheet plus a `Filters` sheet (PTC, 2026-09-20). Every
series is a column on that one sheet, so two datasets a user wants to compare sit
side by side on the same row; before this they were on separate sheets, where
comparing them meant copying one across. `analytics_layout` builds the table and the
CSV writer emits the same one.

Three rules this module exists to keep:

- **A missing value is an empty cell.** `None` is appended as `None`, which openpyxl
  writes as a genuinely blank cell — not `0`, not the string `"None"`. A real `0`
  survives as `0`. In a settlement spreadsheet those are different numbers.
- **The Date column holds real dates.** A date written as text cannot be sorted,
  filtered or subtracted in Excel, and these files are opened to be worked in. Write-only
  mode has no styles, so the cells carry their own number format.
- **The directory is created before the first write.** `ANALYTICS_EXPORT_DIR` defaults
  to `/var/lib/epms/exports`, which does not exist on a dev machine and is not
  writable by the service user on the PROD box either. Failing at `makedirs` with the
  path and the setting named beats a bare `PermissionError` out of the zip writer.
"""

import logging
import os
from datetime import date, datetime
from typing import Any, Dict, Iterable, List, Optional, Sequence

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell

from . import analytics_layout
from .analytics_csv import COLUMNS, SUPPORTED_FORMATS, filename_for, filter_pairs

logger = logging.getLogger(__name__)

__all__ = ["ExportWriteError", "DATA_SHEET", "FILTERS_SHEET", "write_xlsx",
           "filename_for", "COLUMNS", "SUPPORTED_FORMATS"]

#: The filters sheet is named in the spec; the frontend links to it by name.
FILTERS_SHEET = "Filters"

#: The single data sheet. One sheet per dataset was replaced on 2026-09-20.
DATA_SHEET = "Data"

#: How the Date column renders. `dd mmm yyyy` is unambiguous to a reader in India
#: and to one in the US, which `dd/mm/yyyy` is not.
DATE_NUMBER_FORMAT = "dd mmm yyyy"

class ExportWriteError(OSError):
    """The export file could not be written — directory, permissions or disk."""


def _ensure_directory(path: str) -> None:
    directory = os.path.dirname(os.path.abspath(path))
    try:
        os.makedirs(directory, exist_ok=True)
    except OSError as exc:
        raise ExportWriteError(
            f"Cannot create the analytics export directory {directory!r}: {exc}. "
            f"Check ANALYTICS_EXPORT_DIR and that the service user owns it."
        ) from exc
    if not os.path.isdir(directory):  # pragma: no cover - makedirs would have raised
        raise ExportWriteError(
            f"Analytics export directory {directory!r} is not a directory. "
            f"Check ANALYTICS_EXPORT_DIR.")
    if not os.access(directory, os.W_OK):
        raise ExportWriteError(
            f"Analytics export directory {directory!r} is not writable by "
            f"{os.getuid()}. Check ANALYTICS_EXPORT_DIR and its ownership.")


def _date_cell(sheet, value: date) -> Any:
    """A real date cell, formatted. Write-only mode has no styles to inherit."""
    cell = WriteOnlyCell(sheet, value=value)
    cell.number_format = DATE_NUMBER_FORMAT
    return cell


def _cell(sheet, value: Any) -> Any:
    """`None` stays `None` so openpyxl writes a blank cell. `0` stays `0`."""
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return _date_cell(sheet, value)
    return value


def write_xlsx(response: Dict[str, Any], filters: Optional[Dict[str, Any]],
               path: str, on_row=None) -> Dict[str, Any]:
    """Write the `Data` sheet plus `Filters`; return `{path, bytes, row_count}`.

    `row_count` counts data rows only — the heading row and the `Filters` sheet are
    not part of what the user asked for, and the history row reports this number back
    to them as "rows exported".
    """
    response = response or {}

    _ensure_directory(path)

    workbook = Workbook(write_only=True)
    row_count = 0
    try:
        sheet = workbook.create_sheet(title=DATA_SHEET)
        sheet.append(analytics_layout.column_titles(response, filters))
        # `data_rows` is a generator and is kept one: a year of individual blocks is
        # 35,040 rows wide enough to matter, and write-only mode streams each away.
        # `on_row` is how a background export reports its progress.
        for row in analytics_layout.data_rows(response, filters, on_row=on_row):
            sheet.append([_cell(sheet, value) for value in row])
            row_count += 1

        filters_sheet = workbook.create_sheet(title=FILTERS_SHEET)
        filters_sheet.append(["Filter", "Value"])
        for label, value in filter_pairs(filters, response):
            filters_sheet.append([label, value])

        workbook.save(path)
    except ExportWriteError:
        raise
    except OSError as exc:
        raise ExportWriteError(
            f"Cannot write the analytics export {path!r}: {exc}. "
            f"Check ANALYTICS_EXPORT_DIR, its free space and its permissions."
        ) from exc
    finally:
        workbook.close()

    size = os.path.getsize(path)
    logger.info("Analytics XLSX export written: path=%s bytes=%s rows=%s columns=%s",
                path, size, row_count,
                len(analytics_layout.column_titles(response, filters)))
    return {"path": path, "bytes": size, "row_count": row_count}
