"""
The Shift Status Report download, laid out like EDP's own shift sheet.

EDP's sheet is: a "Shift Status Report" title, a "Shift Person Name" row, a header
block, then one row per block. The download keeps that layout and repeats it for
every date in the range, one section under the other on a single sheet, with two
additions the report needs:

* a Date column before Time, because a shift that runs past midnight has rows on
  two dates;
* I-DAM Status and RTM Status each split into Bid and Traded under the sheet's own
  heading — EDP's sheet records the bid, and the traded volume beside it is what
  explains "Buy bid not cleared completely".
"""

import io
import re
from datetime import datetime
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import shift_status as rules

SHEET_TITLE = "Shift Status Report"
REMARKS_HEADER = (
    "Remarks\n(Specify Reason if Under Drawl/Over Drawl is More/Less than "
    f"{rules.REMARK_THRESHOLD_MW:g} MW in any Block)")

#: (header, sub-header or None, row key, width). A shared header spans consecutive
#: columns; a column without a sub-header spans both header rows.
COLUMNS = [
    ("Date", None, "date", 12),
    ("Time", "From", "time_from", 9),
    ("Time", "To", "time_to", 9),
    ("Block", None, "time_block", 8),
    ("DAY Ahead Forecast", None, "day_ahead_forecast", 13),
    ("Real Time Expected Forecast", None, "real_time_forecast", 14),
    ("Net Schedule (As per SRLDC)", None, "net_schedule", 14),
    ("Actual Drawal (As per SRLDC)", None, "actual_drawal", 12),
    ("Under Drawl(-)\nOver Drawl (+)", None, "deviation", 13),
    ("I-DAM Status", "Bid", "idam_bid", 10),
    ("I-DAM Status", "Traded", "idam_traded", 10),
    ("RTM Status", "Bid", "rtm_bid", 10),
    ("RTM Status", "Traded", "rtm_traded", 10),
    (REMARKS_HEADER, None, "remark", 55),
]

_THIN = Side(style="thin", color="000000")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_BOLD = Font(bold=True)
_TITLE_FONT = Font(bold=True, size=13)
_FLAG_FILL = PatternFill("solid", fgColor="FCE4D6")


def display_date(iso: str) -> str:
    """`28-09-2026`, the day-first form every EPMS download uses."""
    return datetime.strptime(iso, "%Y-%m-%d").strftime("%d-%m-%Y")


def filename_for(from_date: str, to_date: str, shift_name: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "-", shift_name or "").strip("-") or "Shift"
    span = from_date if from_date == to_date else f"{from_date}-to-{to_date}"
    return f"Shift-Status-Report-{span}-{slug}.xlsx"


def _header_spans() -> List[tuple]:
    """`(first_col, last_col, header, has_sub)` for each run of equal headers."""
    spans, col = [], 1
    while col <= len(COLUMNS):
        header, sub = COLUMNS[col - 1][0], COLUMNS[col - 1][1]
        last = col
        while last < len(COLUMNS) and COLUMNS[last][0] == header:
            last += 1
        spans.append((col, last, header, sub is not None))
        col = last + 1
    return spans


def _box(ws, first_row: int, first_col: int, last_row: int, last_col: int) -> None:
    for row in ws.iter_rows(min_row=first_row, max_row=last_row,
                            min_col=first_col, max_col=last_col):
        for cell in row:
            cell.border = _BORDER


def _write_section(ws, top: int, section: Dict[str, Any], shift: Dict[str, Any]) -> int:
    """Write one date's section starting at row `top`; return the next free row."""
    width = len(COLUMNS)
    last_col = get_column_letter(width)

    title = f"{SHEET_TITLE} — {display_date(section['date'])} · {shift.get('shift_name')}"
    ws.merge_cells(f"A{top}:{last_col}{top}")
    ws[f"A{top}"] = title
    ws[f"A{top}"].font = _TITLE_FONT
    ws[f"A{top}"].alignment = _CENTER
    _box(ws, top, 1, top, width)

    person = top + 1
    ws.merge_cells(f"A{person}:D{person}")
    ws[f"A{person}"] = "Shift Person Name"
    ws[f"A{person}"].font = _BOLD
    ws[f"A{person}"].alignment = _CENTER
    ws.merge_cells(f"E{person}:{last_col}{person}")
    ws[f"E{person}"] = section.get("shift_person") or ""
    ws[f"E{person}"].alignment = _CENTER
    _box(ws, person, 1, person, width)

    head, sub = person + 1, person + 2
    for first, last, header, has_sub in _header_spans():
        a, b = get_column_letter(first), get_column_letter(last)
        if has_sub:
            ws.merge_cells(f"{a}{head}:{b}{head}")
            for col in range(first, last + 1):
                cell = ws.cell(row=sub, column=col, value=COLUMNS[col - 1][1])
                cell.font, cell.alignment = _BOLD, _CENTER
        else:
            ws.merge_cells(f"{a}{head}:{b}{sub}")
        cell = ws.cell(row=head, column=first, value=header)
        cell.font, cell.alignment = _BOLD, _CENTER
    ws.row_dimensions[head].height = 32
    ws.row_dimensions[sub].height = 32
    _box(ws, head, 1, sub, width)

    row = sub + 1
    for block in section.get("blocks") or []:
        for col, (_h, _s, key, _w) in enumerate(COLUMNS, start=1):
            value = block.get(key)
            if key == "date":
                value = display_date(value)
            if key == "remark":
                value = value or None
            cell = ws.cell(row=row, column=col, value=value)
            cell.alignment = _CENTER
            cell.border = _BORDER
            if key == "deviation" and value is not None \
                    and abs(value) > rules.REMARK_THRESHOLD_MW:
                cell.fill = _FLAG_FILL
        row += 1
    return row + 1  # one blank row between sections


def build_workbook(report: Dict[str, Any]) -> bytes:
    """The whole download as xlsx bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE
    for col, (_h, _s, _k, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    shift = report.get("shift") or {}
    top = 1
    for section in report.get("sections") or []:
        top = _write_section(ws, top, section, shift)
    if not report.get("sections"):
        ws["A1"] = "No dates in range"

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
